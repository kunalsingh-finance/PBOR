from __future__ import annotations

import json
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "outputs" / "2026-01"
DB_PATH = PROJECT_ROOT / "pbor_lite.db"

REQUIRED_FILES = [
    "daily_returns.csv",
    "monthly_returns.csv",
    "attribution.csv",
    "attribution_reconciliation.csv",
    "breaks.csv",
    "qa_ingest_summary.csv",
    "recon_exceptions.csv",
    "signoff_summary.csv",
    "report.xlsx",
    "summary.json",
    "onepager.md",
    "onepager.pdf",
    "tearsheet.png",
    "controls_table.png",
]

REQUIRED_SHEETS = [
    "Summary",
    "MonthlyReturns",
    "DailyReturns",
    "Attribution",
    "AttrReconciliation",
    "Breaks",
    "IngestQA",
    "AutoReconExceptions",
    "SignOffSummary",
]

REQUIRED_SUMMARY_KEYS = [
    "total_recon_records",
    "matched_recon_records",
    "open_recon_exceptions",
    "high_severity_recon_exceptions",
    "cash_break_amount",
    "market_value_break_amount",
    "recon_match_rate",
    "ready_for_signoff",
    "signoff_ready",
    "failed_control_areas",
    "open_high_severity_recon_exceptions",
    "open_high_severity_qa_breaks",
    "sla_breached_recon_exceptions",
    "watchlist_recon_exceptions",
]

REQUIRED_RECON_COLUMNS = [
    "asof_date",
    "portfolio_id",
    "account_id",
    "security_id",
    "ticker",
    "currency",
    "break_type",
    "status",
    "workflow_status",
    "severity",
    "sla_bucket",
    "root_cause",
    "resolution_note",
    "action_required",
    "owner",
    "age_days",
]

REQUIRED_SIGNOFF_ROWS = [
    "Attribution Reconciliation",
    "QA Controls",
    "PBOR vs Custodian Positions",
    "Cash Reconciliation",
    "Final Reporting Sign-Off",
]

REQUIRED_RECON_STATUSES = {
    "QUANTITY_BREAK",
    "PRICE_BREAK",
    "MISSING_IN_CUSTODIAN",
    "MISSING_IN_INTERNAL",
    "CASH_BREAK",
    "MATCHED",
}

REQUIRED_DB_TABLES = [
    "pbor_daily_positions",
    "pbor_daily_returns",
    "pbor_monthly_returns",
    "pbor_attribution_monthly",
    "pbor_breaks",
    "pbor_recon_exceptions",
    "pbor_signoff_summary",
]


class Verifier:
    def __init__(self) -> None:
        self.failures: list[str] = []

    def check(self, condition: bool, message: str) -> None:
        if condition:
            print(f"[PASS] {message}")
        else:
            print(f"[FAIL] {message}")
            self.failures.append(message)

    def finish(self) -> None:
        if self.failures:
            raise SystemExit(1)
        print("Demo verification passed.")


def main() -> None:
    verifier = Verifier()

    verifier.check(OUTPUT_DIR.exists(), "outputs/2026-01 exists")
    for file_name in REQUIRED_FILES:
        verifier.check((OUTPUT_DIR / file_name).exists(), f"{file_name} exists")

    workbook_path = OUTPUT_DIR / "report.xlsx"
    if workbook_path.exists():
        workbook = load_workbook(workbook_path, read_only=True)
        for sheet_name in REQUIRED_SHEETS:
            verifier.check(sheet_name in workbook.sheetnames, f"report.xlsx includes {sheet_name}")
    else:
        for sheet_name in REQUIRED_SHEETS:
            verifier.check(False, f"report.xlsx includes {sheet_name}")

    summary_path = OUTPUT_DIR / "summary.json"
    if summary_path.exists():
        summary = json.loads(summary_path.read_text(encoding="utf-8"))
        for key in REQUIRED_SUMMARY_KEYS:
            verifier.check(key in summary, f"summary.json includes {key}")
    else:
        summary = {}
        for key in REQUIRED_SUMMARY_KEYS:
            verifier.check(False, f"summary.json includes {key}")

    recon_path = OUTPUT_DIR / "recon_exceptions.csv"
    if recon_path.exists():
        recon = pd.read_csv(recon_path)
        for column in REQUIRED_RECON_COLUMNS:
            verifier.check(column in recon.columns, f"recon_exceptions.csv includes {column}")
        statuses = set(recon["status"].astype(str)) if "status" in recon.columns else set()
        for status in sorted(REQUIRED_RECON_STATUSES):
            verifier.check(status in statuses, f"recon_exceptions.csv includes {status}")
    else:
        recon = pd.DataFrame()
        for column in REQUIRED_RECON_COLUMNS:
            verifier.check(False, f"recon_exceptions.csv includes {column}")
        for status in sorted(REQUIRED_RECON_STATUSES):
            verifier.check(False, f"recon_exceptions.csv includes {status}")

    signoff_path = OUTPUT_DIR / "signoff_summary.csv"
    if signoff_path.exists():
        signoff = pd.read_csv(signoff_path)
        rows = set(signoff["control_area"].astype(str)) if "control_area" in signoff.columns else set()
        for control_area in REQUIRED_SIGNOFF_ROWS:
            verifier.check(control_area in rows, f"signoff_summary.csv includes {control_area}")
    else:
        for control_area in REQUIRED_SIGNOFF_ROWS:
            verifier.check(False, f"signoff_summary.csv includes {control_area}")

    verifier.check(DB_PATH.exists(), "pbor_lite.db exists")
    if DB_PATH.exists():
        conn = sqlite3.connect(DB_PATH)
        try:
            tables = {
                row[0]
                for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")
            }
            for table_name in REQUIRED_DB_TABLES:
                verifier.check(table_name in tables, f"SQLite includes {table_name}")

            recon_columns = {
                row[1] for row in conn.execute("PRAGMA table_info(pbor_recon_exceptions)")
            } if "pbor_recon_exceptions" in tables else set()
            for column in ["workflow_status", "sla_bucket", "action_required"]:
                verifier.check(column in recon_columns, f"pbor_recon_exceptions includes {column}")

            if "pbor_signoff_summary" in tables:
                signoff_rows = {
                    row[0]
                    for row in conn.execute("SELECT control_area FROM pbor_signoff_summary")
                }
                for control_area in REQUIRED_SIGNOFF_ROWS:
                    verifier.check(control_area in signoff_rows, f"SQLite signoff includes {control_area}")
                signoff_count = conn.execute("SELECT COUNT(*) FROM pbor_signoff_summary").fetchone()[0]
                verifier.check(signoff_count == 5, "pbor_signoff_summary has exactly 5 rows")
            else:
                for control_area in REQUIRED_SIGNOFF_ROWS:
                    verifier.check(False, f"SQLite signoff includes {control_area}")
                verifier.check(False, "pbor_signoff_summary has exactly 5 rows")

            if "pbor_recon_exceptions" in tables:
                recon_count = conn.execute("SELECT COUNT(*) FROM pbor_recon_exceptions").fetchone()[0]
                verifier.check(recon_count > 0, "pbor_recon_exceptions has rows after recon demo run")
            else:
                verifier.check(False, "pbor_recon_exceptions has rows after recon demo run")
        finally:
            conn.close()
    else:
        for table_name in REQUIRED_DB_TABLES:
            verifier.check(False, f"SQLite includes {table_name}")
        for column in ["workflow_status", "sla_bucket", "action_required"]:
            verifier.check(False, f"pbor_recon_exceptions includes {column}")
        for control_area in REQUIRED_SIGNOFF_ROWS:
            verifier.check(False, f"SQLite signoff includes {control_area}")
        verifier.check(False, "pbor_signoff_summary has exactly 5 rows")
        verifier.check(False, "pbor_recon_exceptions has rows after recon demo run")

    verifier.finish()


if __name__ == "__main__":
    main()
